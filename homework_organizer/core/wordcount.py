"""字数限制 / 作业检查：统计学生文档字数并按检索信息汇总检查结果。

本模块扫描指定文档，统计其字数，并判断是否低于设定阈值。低于阈值的文件
由图形界面标记为红色，便于快速识别需要退回的作业。

除逐文件的字数明细外，本模块还支持"作业检查"汇总：导入一份**检索信息表**
（每行一条，可写姓名、学号、作业名等任意内容），将文件名按分隔符切成任意
多段后与每条信息做**完全匹配**，再统计每条信息对应的识别份数与字数达标情况。

处理流程分为三个独立环节，便于分别定位问题：

    1. :func:`extract_text`        从不同格式文档中提取纯文本
    2. :func:`count_words`         按统一口径统计字数
    3. :func:`apply_threshold`     判定是否低于阈值

"提取失败"与"字数不足"被区分处理：文档无法读取时字数记为未知，不计入低于
阈值的清单，避免将"文件损坏"误判为"内容不足"。汇总统计同样遵循这一原则：
读取失败的份数单独成列，既不并入未达标，也不影响该行的绿标判定。

字数统计口径：
    - 每个中文字符计 1 字。
    - 每个英文单词计 1 字（不以字母计数）。
    - 标点符号与空白字符不计入。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Sequence

__all__ = [
    "WordCountResult",
    "HomeworkCheckRow",
    "HomeworkCheckSummary",
    "SAMPLE_CHECK_INFO_TEXT",
    "SUPPORTED_EXTENSIONS",
    "extract_text",
    "count_words",
    "count_words_in_file",
    "apply_threshold",
    "scan_documents_wordcount",
    "classify_file_by_info",
    "build_homework_check",
    "format_homework_check_summary",
    "build_wordcount_report",
]

#: 支持提取文本的扩展名集合，与 :func:`extract_text` 的实现范围保持一致。
SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(
    {".txt", ".md", ".docx", ".pdf", ".xlsx", ".xlsm"}
)

#: 检索信息表的示例内容，供图形界面的「填入示例」按钮与命令行 ``--show-info-sample`` 使用。
SAMPLE_CHECK_INFO_TEXT = """# 每行一条检索信息，与文件名的每一段做完全匹配
# 可以写姓名、学号、作业名等任意内容
# 以 # 开头的是注释行，空行会被忽略
张三
李四
王五
20230101
第一次作业"""

# 中文字符匹配模式，涵盖基本汉字区与扩展 A 区。
_CJK_PATTERN = re.compile(r"[\u4e00-\u9fff\u3400-\u4dbf]")

# 英文单词匹配模式，允许单词内部包含连字符与撇号（如 don't、well-known）。
_WORD_PATTERN = re.compile(r"[A-Za-z0-9_]+(?:['\-][A-Za-z0-9_]+)*")

#: 中文全角字符的显示宽度按 2 计，用于在等宽字体环境下对齐文本表格。
_CJK_WIDTH_PATTERN = re.compile(r"[\u1100-\u115f\u2e80-\ua4cf\ua960-\ua97f"
                                r"\uac00-\ud7a3\uf900-\ufaff\ufe10-\ufe19"
                                r"\ufe30-\ufe6b\uff00-\uff60\uffe0-\uffe6]")


@dataclass
class WordCountResult:
    """单个文档的字数统计结果。"""

    path: Path
    """文档路径。"""

    count: int | None
    """统计所得字数；``None`` 表示统计失败。"""

    below_threshold: bool
    """是否低于设定阈值。统计失败时恒为 ``False``。"""

    error: str = ""
    """统计失败的原因；成功时为空字符串。"""

    @property
    def is_ok(self) -> bool:
        """统计是否成功完成。

        Returns:
            成功统计出字数时返回 ``True``。

        Note:
            该属性用于区分"无法读取"与"字数不足"两种状态。
        """
        return self.count is not None and not self.error


@dataclass
class HomeworkCheckRow:
    """一条检索信息的汇总结果。

    对应汇总表中的一行：某一个关键词命中了哪些文件、其中多少份字数达标。
    """

    keyword: str
    """检索信息原文（已去除首尾空白）。"""

    matched: list[WordCountResult] = field(default_factory=list)
    """文件名命中该信息的全部统计结果，按扫描顺序排列。"""

    @property
    def total(self) -> int:
        """识别到的份数（含读取失败的部分）。

        Returns:
            命中该信息的文件总数。
        """
        return len(self.matched)

    @property
    def failed(self) -> int:
        """字数读取失败的份数。

        这类文件（如扫描件 PDF）单独统计，既不并入未达标，也不影响绿标判定。

        Returns:
            统计失败的文件数量。
        """
        return sum(1 for result in self.matched if not result.is_ok)

    @property
    def passed(self) -> int:
        """字数达标的份数。

        Returns:
            统计成功且不低于阈值的文件数量。
        """
        return sum(
            1 for result in self.matched if result.is_ok and not result.below_threshold
        )

    @property
    def below(self) -> int:
        """字数未达标的份数。

        Returns:
            统计成功但低于阈值的文件数量。
        """
        return sum(
            1 for result in self.matched if result.is_ok and result.below_threshold
        )

    @property
    def is_green(self) -> bool:
        """该行是否标绿。

        判定条件为"字数达标份数 ≥ 1"：只要有一份达标即通过，哪怕同时存在
        未达标的文件（此时未达标份数会体现在对应列里）。

        Returns:
            达标份数大于等于 1 时返回 ``True``。
        """
        return self.passed >= 1

    @property
    def matched_names(self) -> list[str]:
        """命中该信息的文件名清单。

        Returns:
            文件名字符串列表，供报告与提示文案使用。
        """
        return [result.path.name for result in self.matched]


@dataclass
class HomeworkCheckSummary:
    """作业检查的整体汇总结果。

    由 :func:`build_homework_check` 产出，供图形界面渲染汇总表、命令行生成
    报告，以及判定是否给出"总绿勾"。
    """

    rows: list[HomeworkCheckRow] = field(default_factory=list)
    """按检索信息顺序排列的汇总行。"""

    threshold: int = 0
    """本次检查使用的字数下限。"""

    unmatched: list[str] = field(default_factory=list)
    """文件名未命中任何检索信息的文件清单（不参与任何一行统计）。"""

    @property
    def all_green(self) -> bool:
        """是否全部通过。

        Returns:
            存在至少一行且所有行都标绿时返回 ``True``。

        Note:
            空表刻意返回 ``False``：一条信息都没填时若给出绿勾，会误导用户
            以为检查已经完成。
        """
        return bool(self.rows) and all(row.is_green for row in self.rows)

    @property
    def red_rows(self) -> list[HomeworkCheckRow]:
        """标红的行（达标份数为 0）。

        Returns:
            未通过的行列表，供弹窗提示点名使用。
        """
        return [row for row in self.rows if not row.is_green]


def extract_text(path: str | Path) -> str:
    """从文档中提取纯文本内容。

    支持的格式：``.txt``、``.md``、``.docx``、``.pdf``、``.xlsx``、``.xlsm``。

    Args:
        path: 文档路径。

    Returns:
        文档的纯文本内容，段落之间以换行符连接。

    Raises:
        ValueError: 文件格式不受支持。``.doc`` 与 ``.docx`` 为不同格式，
            ``python-docx`` 无法读取 ``.doc``，故一并归入不支持范围。
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in (".txt", ".md"):
        # utf-8-sig 可自动去除 Windows 记事本写入的 BOM。
        return path.read_text(encoding="utf-8-sig", errors="replace")

    if suffix == ".docx":
        return _extract_docx(path)

    if suffix == ".pdf":
        return _extract_pdf(path)

    if suffix in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)

    raise ValueError(f"不支持的文档格式：{suffix or '（无扩展名）'}")


def _extract_docx(path: Path) -> str:
    """提取 ``.docx`` 的正文段落与表格内容。

    作业模板常将内容填入表格，仅读取段落会遗漏大量文字并导致字数被低估。

    Args:
        path: 文档路径。

    Returns:
        段落与表格文本以换行符连接的结果。
    """
    # 延迟导入，使未安装对应依赖时仅在实际处理该格式时报错。
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    return "\n".join(parts)


def _extract_pdf(path: Path) -> str:
    """逐页提取 ``.pdf`` 中的文本。

    图片型扫描件不含文字层，提取结果为空字符串，此为 PDF 格式的固有局限，
    需要 OCR 才能处理。

    Args:
        path: 文档路径。

    Returns:
        各页文本以换行符连接的结果。
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []

    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # 单页提取失败不影响其余页面的处理。
            continue

    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    """提取 ``.xlsx`` 中所有工作表的单元格内容。

    仅取单元格的显示值，不包含公式文本。

    Args:
        path: 文档路径。

    Returns:
        所有单元格值以换行符连接的结果。
    """
    from openpyxl import load_workbook

    # data_only 读取公式计算结果；read_only 降低大文件的内存占用。
    workbook = load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is None:
                    continue
                parts.append(str(value))

    # 只读模式需显式关闭以释放文件句柄，否则 Windows 上文件会持续被占用。
    workbook.close()
    return "\n".join(parts)


def count_words(text: str) -> int:
    """按统一口径统计文本字数。

    中文字符与英文单词分别统计后相加。两个匹配模式的字符集互不重叠，
    不存在重复计数。

    Args:
        text: 待统计的文本。

    Returns:
        统计所得字数；空文本返回 ``0``。
    """
    if not text:
        return 0

    chinese_count = len(_CJK_PATTERN.findall(text))
    english_count = len(_WORD_PATTERN.findall(text))
    return chinese_count + english_count


def count_words_in_file(path: str | Path) -> WordCountResult:
    """读取文档并统计字数。

    本函数不抛出异常：批量处理时单个文件损坏不应中断整体流程，所有异常均
    转换为带原因的失败结果。

    Args:
        path: 文档路径。

    Returns:
        统计结果。失败时 ``count`` 为 ``None``，``error`` 说明原因。
    """
    path = Path(path)

    try:
        text = extract_text(path)
    except ValueError as exc:
        return WordCountResult(
            path=path, count=None, below_threshold=False, error=str(exc)
        )
    except Exception as exc:
        # 附带异常类型名，因为部分异常的字符串表示为空。
        return WordCountResult(
            path=path,
            count=None,
            below_threshold=False,
            error=f"读取失败：{type(exc).__name__}: {exc}",
        )

    return WordCountResult(
        path=path,
        count=count_words(text),
        below_threshold=False,
    )


def apply_threshold(
    results: Sequence[WordCountResult],
    threshold: int,
) -> list[WordCountResult]:
    """标记各统计结果是否低于阈值。

    本函数就地修改传入对象的 ``below_threshold`` 字段。判定采用严格小于：
    恰好等于阈值视为合格。统计失败的结果不标记为低于阈值。

    Args:
        results: 统计结果序列。
        threshold: 字数下限。

    Returns:
        已标记完成的同一列表。
    """
    for result in results:
        if result.count is None:
            result.below_threshold = False
            continue
        result.below_threshold = result.count < threshold

    return results


def scan_documents_wordcount(
    records: Sequence,
    threshold: int,
    progress: Callable[[int, int, str], None] | None = None,
    only_supported: bool = True,
) -> tuple[list[WordCountResult], list[tuple[str, str]]]:
    """批量统计文档字数并标记低于阈值的文件。

    Args:
        records: 扫描记录序列，需具有 ``path`` 与 ``name`` 属性。
        threshold: 字数下限。
        progress: 进度回调，签名为 ``(当前序号, 总数, 说明文本)``。
        only_supported: 为 ``True`` 时不统计不受支持的格式，但仍将其记入
            跳过清单，以便用户核对数量。

    Returns:
        二元组 ``(统计结果列表, 跳过清单)``。结果列表已应用阈值标记，且
        包含统计成功与失败的完整记录，便于报告展示全貌。
    """
    results: list[WordCountResult] = []
    skipped: list[tuple[str, str]] = []

    if only_supported:
        targets = [
            record
            for record in records
            if record.path.suffix.lower() in SUPPORTED_EXTENSIONS
        ]
        for record in records:
            if record.path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                skipped.append(
                    (
                        str(record.path),
                        f"不支持的格式：{record.path.suffix or '（无扩展名）'}",
                    )
                )
    else:
        targets = list(records)

    total = len(targets)

    for index, record in enumerate(targets, start=1):
        if progress is not None:
            progress(index, total, record.name)

        result = count_words_in_file(record.path)
        if not result.is_ok:
            skipped.append((str(record.path), result.error))
        results.append(result)

    apply_threshold(results, threshold)
    return results, skipped


def classify_file_by_info(
    result: WordCountResult,
    keywords: Sequence[str],
) -> list[str]:
    """找出某个文件的文件名命中了哪些检索信息。

    匹配口径与"作业归类"完全一致：文件名按 ``_``、``+``、空格切成任意多段，
    某一段与检索信息**整段相等**才算命中。因此 ``"张三"`` 不会命中
    ``"张三丰"``——这是刻意的，避免子串包含带来误统计。

    Args:
        result: 单个文件的字数统计结果。
        keywords: 检索信息序列。

    Returns:
        命中的检索信息列表，按 ``keywords`` 原顺序排列；一个文件可命中多条。
    """
    # 延迟导入打破循环依赖：rename 不依赖本模块，本模块只在需要时才取它。
    from . import rename as rename_core

    if not keywords:
        return []

    stem = result.path.stem
    segments = set(rename_core.split_segments(stem))
    return [keyword for keyword in keywords if keyword in segments]


def build_homework_check(
    results: Sequence[WordCountResult],
    keywords: Sequence[str],
    threshold: int,
) -> HomeworkCheckSummary:
    """按检索信息汇总检查结果。

    逐个文件比对，把同一份统计结果分发给它命中的每一条检索信息——因此一个
    文件名里同时含姓名与学号时，会同时计入这两个关键词所在的行。

    Args:
        results: 已应用阈值标记的字数统计结果序列。
        keywords: 检索信息序列，汇总行的顺序与之保持一致。
        threshold: 本次检查使用的字数下限，仅用于记录。

    Returns:
        汇总结果。``unmatched`` 收集了未命中任何检索信息的文件名。
    """
    # 先建好各行，保证顺序与检索信息一致（即使某条信息一份都没命中）。
    rows_by_keyword: dict[str, HomeworkCheckRow] = {
        keyword: HomeworkCheckRow(keyword=keyword) for keyword in keywords
    }
    unmatched: list[str] = []

    for result in results:
        hits = classify_file_by_info(result, keywords)
        if not hits:
            unmatched.append(result.path.name)
            continue
        for keyword in hits:
            rows_by_keyword[keyword].matched.append(result)

    return HomeworkCheckSummary(
        rows=[rows_by_keyword[keyword] for keyword in keywords],
        threshold=threshold,
        unmatched=unmatched,
    )


def _display_width(text: str) -> int:
    """计算文本在等宽终端下的显示宽度。

    中日韩全角字符占 2 个字符宽，其余占 1 个。

    Args:
        text: 待计算宽度的文本。

    Returns:
        显示宽度（列数）。
    """
    return sum(2 if _CJK_WIDTH_PATTERN.match(char) else 1 for char in text)


def _pad(text: str, width: int) -> str:
    """把文本右侧补齐到指定显示宽度。

    Args:
        text: 待补齐的文本。
        width: 目标显示宽度。

    Returns:
        补齐后的文本；已超宽时原样返回。
    """
    padding = width - _display_width(text)
    return text + " " * padding if padding > 0 else text


def format_homework_check_summary(summary: HomeworkCheckSummary) -> str:
    """把汇总结果渲染为等宽文本表格。

    Args:
        summary: 汇总结果。

    Returns:
        多行文本表格；``summary.rows`` 为空时返回一行说明。
    """
    if not summary.rows:
        return "（未填写检索信息，本次仅输出逐文件明细）"

    #: 表格各列的标题，与图形界面的汇总表保持一致。
    headers = ["输入信息", "识别份数", "达标份数", "未达标份数", "读取失败"]
    # 每列的显示宽度取"表头宽"与"该列所有单元格最大宽"中的较大者。
    widths = [_display_width(header) for header in headers]

    cells: list[list[str]] = []
    for row in summary.rows:
        row_cells = [
            row.keyword,
            str(row.total),
            str(row.passed),
            str(row.below),
            str(row.failed),
        ]
        cells.append(row_cells)
        for index, cell in enumerate(row_cells):
            widths[index] = max(widths[index], _display_width(cell))

    lines: list[str] = []
    lines.append("【按检索信息汇总】")
    lines.append(
        "  "
        + "  ".join(_pad(header, widths[i]) for i, header in enumerate(headers))
    )
    lines.append("  " + "  ".join("-" * width for width in widths))

    for row, row_cells in zip(summary.rows, cells):
        line = "  " + "  ".join(
            _pad(cell, widths[i]) for i, cell in enumerate(row_cells)
        )
        # 标红行加一个醒目前缀，方便在纯文本里一眼扫到。
        if not row.is_green:
            line += "  <-- 未通过"
        lines.append(line)

    lines.append("")
    if summary.all_green:
        lines.append("  ✓ 全部通过（每条检索信息都至少有 1 份达标）")
    else:
        red_names = "、".join(row.keyword for row in summary.red_rows)
        lines.append(f"  ⚠ 未通过：{red_names}")

    if summary.unmatched:
        lines.append("")
        lines.append(f"  未被任何检索信息识别到的文件（{len(summary.unmatched)} 个）：")
        for name in summary.unmatched[:20]:
            lines.append(f"    - {name}")
        if len(summary.unmatched) > 20:
            lines.append(f"    ...（其余 {len(summary.unmatched) - 20} 个已省略）")

    return "\n".join(lines)


def build_wordcount_report(
    results: Sequence[WordCountResult],
    threshold: int,
    root: str | Path | None = None,
    summary: HomeworkCheckSummary | None = None,
) -> str:
    """生成作业检查报告。

    报告分为按检索信息汇总、低于要求、统计失败、字数达标四个区块
    （未提供 ``summary`` 时省略第一个区块）。低于要求的文件列在靠前位置，
    并给出与阈值之间的差值。

    Args:
        results: 已应用阈值的统计结果序列。
        threshold: 字数下限。
        root: 可选的检查目录，用于在报告头部展示。
        summary: 可选的按检索信息汇总结果；为 ``None`` 时不输出汇总区块。

    Returns:
        多行报告文本。
    """
    lines: list[str] = []

    below = [r for r in results if r.is_ok and r.below_threshold]
    failed = [r for r in results if not r.is_ok]
    passed = [r for r in results if r.is_ok and not r.below_threshold]

    lines.append("=" * 70)
    lines.append("           作业检查报告")
    lines.append("=" * 70)
    lines.append(f"检查时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"字数要求：不低于 {threshold} 字")
    if root is not None:
        lines.append(f"检查目录：{Path(root)}")

    # 汇总区块放在最前：它回答的是"我关心的每一条信息情况如何"。
    if summary is not None:
        lines.append("")
        lines.append("-" * 70)
        lines.append(
            f"【按检索信息汇总 · 共 {len(summary.rows)} 条】"
            "达标份数 ≥ 1 即通过"
        )
        lines.append("-" * 70)
        lines.append(format_homework_check_summary(summary))

    lines.append("")
    lines.append("-" * 70)
    lines.append(f"【低于字数要求 · 共 {len(below)} 人】需要重点关注")
    lines.append("-" * 70)
    if below:
        for i, result in enumerate(below, start=1):
            gap = threshold - (result.count or 0)
            lines.append(f"  {i:>3}. {result.path.name}")
            lines.append(f"       实际字数：{result.count}  |  还差 {gap} 字")
    else:
        lines.append("  （全部达标）")

    lines.append("")
    lines.append("-" * 70)
    lines.append(f"【统计失败 · 共 {len(failed)} 个】需人工检查")
    lines.append("-" * 70)
    if failed:
        for i, result in enumerate(failed, start=1):
            lines.append(f"  {i:>3}. {result.path.name}")
            lines.append(f"       原因：{result.error}")
    else:
        lines.append("  （无）")

    lines.append("")
    lines.append("-" * 70)
    lines.append(f"【字数达标 · 共 {len(passed)} 个】")
    lines.append("-" * 70)
    if passed:
        for i, result in enumerate(passed, start=1):
            lines.append(f"  {i:>3}. {result.path.name}  （{result.count} 字）")
    else:
        lines.append("  （无）")

    lines.append("")
    lines.append("=" * 70)
    lines.append(
        f"汇总：检查 {len(results)} 个文件 | "
        f"达标 {len(passed)} 个 | "
        f"低于要求 {len(below)} 个 | "
        f"统计失败 {len(failed)} 个"
    )
    lines.append("=" * 70)

    return "\n".join(lines)
